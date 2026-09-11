"""Read annotator-filled .xlsx workbooks back into header-keyed rows.

Filled calibration workbooks come back through the Drive sync mount, which
sometimes produces *streamed* zips whose central directory never synced — the
file opens fine in Google Sheets but ``openpyxl`` refuses it. The reader
falls back to walking the local-file headers of the zip directly, which is
enough to recover the first worksheet and the shared-string table.
"""

from __future__ import annotations

import re
import struct
import xml.etree.ElementTree as ET
import zlib
from pathlib import Path

from loguru import logger

_NS = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"


def _entries_from_streamed_zip(raw: bytes) -> dict[str, bytes]:
    """Extract a streamed zip missing its central directory (Drive artifact)."""
    entries: dict[str, bytes] = {}
    for match in re.finditer(rb"PK\x03\x04", raw):
        off = match.start()
        if off + 30 > len(raw):
            continue
        _, _, _, method, _, _, _, csize, _, nlen, elen = struct.unpack(
            "<IHHHHHIIIHH", raw[off : off + 30]
        )
        name = raw[off + 30 : off + 30 + nlen].decode("utf-8", "replace")
        dstart = off + 30 + nlen + elen
        if method == 8:
            # Streamed entries carry csize=0 in the local header (sizes live
            # in the data descriptor), so the stream's own final block is the
            # only reliable terminator; decompressobj stops there and ignores
            # trailing bytes. When csize IS present, slice to it as well.
            data = raw[dstart : dstart + csize] if csize else raw[dstart:]
            try:
                entries[name] = zlib.decompressobj(-15).decompress(data, 50_000_000)
            except zlib.error as exc:
                # A member that will not inflate is dropped, not fatal: the
                # reader only needs sheet1 + sharedStrings, and a missing
                # REQUIRED member fails loudly at the lookup below.
                logger.warning(
                    f"streamed-zip fallback: could not inflate member '{name}': {exc}"
                )
        elif method == 0 and csize:
            entries[name] = raw[dstart : dstart + csize]
    return entries


def _rows_from_sheet_xml(sheet_xml: bytes, sst_xml: bytes | None) -> list[list]:
    shared: list[str] = []
    if sst_xml:
        for si in ET.fromstring(sst_xml).iter(f"{_NS}si"):
            shared.append("".join(t.text or "" for t in si.iter(f"{_NS}t")))
    grid: dict[int, dict[int, str]] = {}
    max_col = 0
    for row in ET.fromstring(sheet_xml).iter(f"{_NS}row"):
        rnum = int(row.get("r"))
        for cell in row:
            col_letters = re.match(r"[A-Z]+", cell.get("r")).group()
            ci = 0
            for ch in col_letters:
                ci = ci * 26 + (ord(ch) - 64)
            value = cell.find(f"{_NS}v")
            if value is None or value.text is None:
                continue
            text = shared[int(value.text)] if cell.get("t") == "s" else value.text
            grid.setdefault(rnum, {})[ci - 1] = text
            max_col = max(max_col, ci)
    return [[grid[rnum].get(i) for i in range(max_col)] for rnum in sorted(grid)]


def read_workbook_rows(path: Path) -> list[list]:
    """First worksheet as list-of-lists; falls back to manual zip parsing.

    Only a zip-shaped failure (the Drive streamed-zip artifact this module
    exists for) triggers the fallback, and it is logged; anything else — a
    missing file, malformed XML — propagates loudly.
    """
    import zipfile

    from openpyxl import load_workbook
    from openpyxl.utils.exceptions import InvalidFileException

    try:
        wb = load_workbook(path, read_only=True, data_only=True)
        rows = [list(r) for r in wb.worksheets[0].iter_rows(values_only=True)]
        wb.close()
        return rows
    except (zipfile.BadZipFile, InvalidFileException) as exc:
        logger.warning(
            f"openpyxl could not open {path} ({exc}); falling back to "
            "streamed-zip recovery"
        )
        entries = _entries_from_streamed_zip(Path(path).read_bytes())
        if "xl/worksheets/sheet1.xml" not in entries:
            raise ValueError(
                f"{path}: streamed-zip recovery found no first worksheet — "
                "the workbook is not a recoverable xlsx"
            ) from exc
        return _rows_from_sheet_xml(
            entries["xl/worksheets/sheet1.xml"], entries.get("xl/sharedStrings.xml")
        )


def header_dicts(rows: list[list]) -> list[dict[str, str]]:
    """Header-keyed dict per non-empty data row (blank headers dropped)."""
    if not rows:
        return []
    header = [str(h) if h is not None else "" for h in rows[0]]
    out: list[dict[str, str]] = []
    for row in rows[1:]:
        if not any(cell not in (None, "") for cell in row):
            continue
        keyed: dict[str, str] = {}
        for i, name in enumerate(header):
            if not name or name == "None":
                continue
            value = row[i] if i < len(row) else None
            keyed[name] = "" if value is None else str(value)
        out.append(keyed)
    return out
