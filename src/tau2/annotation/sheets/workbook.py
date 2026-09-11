# Copyright Sierra
"""Exported sheet families → formatted .xlsx workbooks (dropdowns + hidden ids).

Replaces the old manual "open the CSV in Sheets and add dropdowns by hand"
step with a reproducible build. Everything is driven by the row models:
dropdowns from ``SheetRow.dropdowns()``, hidden columns from
``hidden_labels()``, widths from the ``ColumnSpec``s — the sheet, the export,
and the analyzer's normalizers can never drift apart. Which workbook to build
is decided by the artifact's manifest ``kind`` (never by header sniffing).

Blind-safety rule: the cold workbook's rubric tab shows each factor's nuance
and ``native_does`` ONLY — naming the typical non-native slip
(``non_native_ref``) would prime the annotator, so it is never rendered.

Publishing: the Sheets import MCP tooling cannot ingest a local .xlsx with its
dropdowns/tabs, so publish by copying the built .xlsx into the Google Drive
desktop mount (``--publish-to``); an uploaded .xlsx opens in Sheets with its
list data-validation showing as native dropdowns.
"""

import shutil
from pathlib import Path
from typing import Optional

from loguru import logger
from openpyxl import Workbook
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from tau2.annotation.artifacts import LoadedArtifact, read_artifact
from tau2.annotation.models import (
    COLD_META_LABELS,
    COLD_TRACE_LABELS,
    COLD_TRANSCRIPT_LABEL,
    COLD_TRANSCRIPT_WIDTH,
    ColdLabel,
    FactorKeyRow,
    JudgePrecisionRow,
)

_DEFAULT_WIDTH = 26
_MAX_DATA_ROW = 1000  # dropdown validation range ceiling
_WRAP_MIN_WIDTH = 40  # columns at least this wide get wrapped text


def _list_validation(options: list[str]) -> DataValidation:
    # Inline (comma-joined) list — option values never contain commas.
    dv = DataValidation(
        type="list", formula1='"' + ",".join(options) + '"', allow_blank=True
    )
    dv.error = "Pick one of the dropdown values."
    dv.promptTitle = "Choose"
    return dv


def _append_text_row(ws, row_idx: int, values: list[str]) -> None:
    """Write one row as EXPLICIT text cells.

    openpyxl infers a leading ``=`` as a formula (data_type 'f'), which would
    let LLM/judge/transcript content execute as a spreadsheet formula when the
    workbook opens. Forcing data_type 's' keeps every cell inert without
    mutating its content.
    """
    for col_idx, value in enumerate(values, start=1):
        cell = ws.cell(row=row_idx, column=col_idx, value=value)
        cell.data_type = "s"


def write_sheet(
    ws,
    headers: list[str],
    rows: list[dict[str, str]],
    *,
    dropdowns: dict[str, list[str]],
    hidden: set[str],
    widths: dict[str, int],
    freeze: str = "A2",
) -> None:
    """Fill one worksheet: text-only cells, dropdowns, widths, wrapping.

    The single place a tau2 annotator worksheet is laid out — every builder
    in this package goes through it, so a column can never be wrapped in one
    sheet and unwrapped in the next.
    """
    _append_text_row(ws, 1, headers)
    for i, row in enumerate(rows, start=2):
        _append_text_row(ws, i, [row.get(h, "") for h in headers])
    ws.freeze_panes = freeze
    wrap = Alignment(wrap_text=True, vertical="top")
    for i, header in enumerate(headers, start=1):
        letter = get_column_letter(i)
        col = ws.column_dimensions[letter]
        col.width = widths.get(header, _DEFAULT_WIDTH)
        if header in hidden:
            col.hidden = True
        if header in dropdowns:
            dv = _list_validation(dropdowns[header])
            ws.add_data_validation(dv)
            dv.add(f"{letter}2:{letter}{_MAX_DATA_ROW}")
        if widths.get(header, 0) >= _WRAP_MIN_WIDTH:
            for row_idx in range(2, len(rows) + 2):
                ws[f"{letter}{row_idx}"].alignment = wrap


def build_precision_workbook(rows: list[dict[str, str]], out: Path) -> Path:
    """One 'Adjudicate' sheet over JudgePrecisionRow columns."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Adjudicate"
    write_sheet(
        ws,
        JudgePrecisionRow.headers(),
        rows,
        dropdowns=JudgePrecisionRow.dropdowns(),
        hidden=JudgePrecisionRow.hidden_labels(),
        widths=JudgePrecisionRow.widths(),
    )
    wb.save(out)
    logger.info(f"wrote {len(rows)} rows -> {out} (Adjudicate)")
    return out


def build_cold_workbook(
    headers: list[str],
    rows: list[dict[str, str]],
    factors_key: list[FactorKeyRow],
    out: Path,
) -> Path:
    """'Conversations' grid (dropdown per factor column) + a blind-safe rubric tab."""
    factor_cols = [h for h in headers if h not in COLD_META_LABELS]
    wb = Workbook()
    grid = wb.active
    grid.title = "Conversations"
    write_sheet(
        grid,
        headers,
        rows,
        dropdowns={c: ColdLabel.options() for c in factor_cols},
        hidden=set(COLD_TRACE_LABELS),
        widths={COLD_TRANSCRIPT_LABEL: COLD_TRANSCRIPT_WIDTH},
    )
    # Blind-safe rubric: nuance + the positive standard only. Naming the
    # typical non-native slip would prime the annotator — non_native_ref is
    # intentionally NEVER rendered here.
    rubric = wb.create_sheet("What each column means")
    rubric_headers = ["What we're checking", "What a native speaker does"]
    rubric_rows = [
        {
            "What we're checking": key.nuance,
            "What a native speaker does": key.native_does,
        }
        for key in factors_key
    ]
    write_sheet(
        rubric,
        rubric_headers,
        rubric_rows,
        dropdowns={},
        hidden=set(),
        widths={rubric_headers[0]: 40, rubric_headers[1]: 70},
    )
    wb.save(out)
    logger.info(
        f"wrote {len(rows)} rows x {len(factor_cols)} factors -> {out} "
        "(Conversations + rubric)"
    )
    return out


def build_workbooks(csv_path: Path) -> list[tuple[Path, str]]:
    """Build the workbook(s) for one exported CSV; returns (xlsx, mode_label)s.

    Mode comes from the manifest kind + the CSV's family role — never from
    header sniffing. Sidecar/key CSVs raise (glob the family's main CSV).
    """
    artifact = read_artifact(Path(csv_path))
    return build_workbooks_for_artifact(artifact, Path(csv_path))


def build_workbooks_for_artifact(
    artifact: LoadedArtifact, csv_path: Path
) -> list[tuple[Path, str]]:
    kind, role = artifact.manifest.kind, artifact.role or "sheet"
    if kind == "nativeness_precision" or (
        kind == "nativeness_cold" and role == "precision"
    ):
        out = csv_path.with_suffix(".xlsx")
        return [(build_precision_workbook(artifact.raw_rows(role), out), "ADJUDICATE")]
    if kind == "nativeness_cold" and role == "sheet":
        raw = artifact.raw_rows("sheet")
        headers = list(raw[0].keys()) if raw else artifact.cold_sheet().headers()
        factors = artifact.rows("factors_key")
        out = csv_path.with_name(csv_path.stem + "_annotate.xlsx")
        return [(build_cold_workbook(headers, raw, factors, out), "ANNOTATE")]
    raise ValueError(
        f"no workbook builder for kind={kind} role={role} — build workbooks "
        "from the annotator-facing sheet of a nativeness family"
    )


def readable_name(src: Path, mode_label: str) -> str:
    """A human, shareable workbook name from the CSV stem + mode, e.g.
    ``es_openai_precision.csv`` + ADJUDICATE → ``es (openai) — ADJUDICATE.xlsx``."""
    stem = src.stem
    for suffix in ("_precision",):
        if stem.endswith(suffix):
            stem = stem[: -len(suffix)]
    parts = stem.split("_")
    pretty = f"{parts[0]} ({'_'.join(parts[1:])})" if len(parts) > 1 else stem
    return f"{pretty} — {mode_label}.xlsx"


def build_and_publish(csv_paths: list[Path], publish_to: Optional[Path] = None) -> int:
    """CLI body for ``tau2 annotate workbook``; returns the number built."""
    if publish_to is not None:
        publish_to.mkdir(parents=True, exist_ok=True)
    built = 0
    for path in csv_paths:
        path = Path(path)
        if not path.exists():
            logger.warning(f"{path} not found; skipping")
            continue
        for out, mode_label in build_workbooks(path):
            built += 1
            if publish_to is not None:
                dest = publish_to / readable_name(path, mode_label)
                shutil.copy2(out, dest)
                logger.info(f"published -> {dest}")
    return built
