# Copyright Sierra
"""Side-by-side localization review sheet: built off the REAL retail arms and
read back from the xlsx, so a column that drifted from what a run executes
fails here rather than in a reviewer's inbox."""

import json

import pytest
from openpyxl import load_workbook
from pydantic import ValidationError

from tau2.annotation.localization_sheet import (
    IN_SUBSET_LABEL,
    TASK_ID_LABEL,
    LocalizationSheetOptions,
    build_localization_sheet,
    source_task_id,
)
from tau2.annotation.models import LocalizationVerdict

LANGS = ["en", "es", "pt", "ko", "zh", "hi"]


@pytest.fixture(scope="module")
def sheet(tmp_path_factory):
    out = tmp_path_factory.mktemp("localization_review")
    xlsx, manifest = build_localization_sheet(
        LocalizationSheetOptions(domain="retail", languages=LANGS, out=out)
    )
    return load_workbook(xlsx), json.loads(manifest.read_text())


def _headers(ws):
    return [c.value for c in ws[1]]


def _column(ws, header):
    return _headers(ws).index(header) + 1


def _text_headers(ws):
    """The scenario-text header of each language group (text/verdict/notes)."""
    return _headers(ws)[2::3]


def _row_for(ws, task_id):
    for row in range(2, ws.max_row + 1):
        if ws.cell(row, 1).value == task_id:
            return row
    raise AssertionError(f"task {task_id} not in the sheet")


def test_one_group_of_three_columns_per_language(sheet):
    wb, _ = sheet
    assert wb.sheetnames == ["Tasks", "How to review"]
    headers = _headers(wb["Tasks"])
    assert headers[:2] == [TASK_ID_LABEL, IN_SUBSET_LABEL]
    # task text / verdict / notes for each of the six languages
    assert len(headers) == 2 + 3 * len(LANGS)
    assert headers[2].startswith("English")
    assert headers[-3].startswith("Hindi")


def test_every_language_column_carries_the_verdict_dropdown(sheet):
    wb, _ = sheet
    ws = wb["Tasks"]
    verdict_cols = [h for h in _headers(ws) if h.endswith("— approve?")]
    assert len(verdict_cols) == len(LANGS)
    formulas = " ".join(dv.formula1 for dv in ws.data_validations.dataValidation)
    for option in LocalizationVerdict.options():
        assert option in formulas
    # One validation range per verdict column, and nothing else validated.
    assert len(ws.data_validations.dataValidation) == len(LANGS)


def test_a_row_holds_each_languages_own_identity(sheet):
    """The point of the sheet: same task, six different callers, side by side."""
    wb, _ = sheet
    ws = wb["Tasks"]
    row = _row_for(ws, "17")
    cells = {
        lang: ws.cell(row, _column(ws, h)).value
        for lang, h in zip(LANGS, _text_headers(ws))
    }
    assert "Fatima Johnson" in cells["en"]
    # ASCII-folded at identity build — the DB and golden actions are
    # unaccented by design (see entity_localization.fold_to_ascii).
    assert "Alba Marin" in cells["es"]
    # Task 17's unit-only change is a destination instruction, not an identity
    # correction, so "Suite 641" stays canonical in every language (see
    # invariants.caller_address_variant).
    assert "Suite 641" in cells["en"]
    assert "Suite 641" in cells["es"]
    assert "Aarti Kumar" in cells["hi"]
    assert "Suite 641" in cells["hi"]
    # Six distinct callers, no cell accidentally copied from the reference.
    assert len({c for c in cells.values()}) == len(LANGS)


def test_shared_benchmark_values_are_identical_across_languages(sheet):
    """Order ids are the shared benchmark and must never be localized."""
    wb, _ = sheet
    ws = wb["Tasks"]
    row = _row_for(ws, "17")
    for header in _text_headers(ws):
        assert "#W8665881" in ws.cell(row, _column(ws, header)).value


def test_run_set_marker_counts_the_canonical_subset(sheet):
    wb, manifest = sheet
    ws = wb["Tasks"]
    marked = sum(1 for row in range(2, ws.max_row + 1) if ws.cell(row, 2).value == "✓")
    assert marked == 50
    assert manifest["canonical_subset"] == "retail_50"
    assert manifest["rows_in_subset"] == marked
    assert manifest["rows"] == ws.max_row - 1


def test_manifest_records_the_task_set_each_column_was_read_from(sheet):
    _, manifest = sheet
    assert manifest["kind"] == "localization_review"
    assert manifest["reference_language"] == "en"
    # The identity arm is what a run executes for every non-English language.
    assert manifest["task_sets"]["en"] == "retail_en"
    assert manifest["task_sets"]["hi"] == "retail_hi_identity"
    assert manifest["git_sha"] is not None


def test_only_subset_keeps_just_the_run_set(tmp_path):
    xlsx, manifest = build_localization_sheet(
        LocalizationSheetOptions(
            domain="retail", languages=LANGS, out=tmp_path, only_subset=True
        )
    )
    ws = load_workbook(xlsx)["Tasks"]
    assert ws.max_row - 1 == 50
    assert all(ws.cell(r, 2).value == "✓" for r in range(2, ws.max_row + 1))
    assert json.loads(manifest.read_text())["only_subset"] is True


@pytest.mark.parametrize(
    "task_id,language,expected",
    [
        ("0_en", "en", "0"),
        ("113_hi_identity", "hi", "113"),
        ("17_zh_identity", "zh", "17"),
    ],
)
def test_source_task_id_strips_both_id_shapes(task_id, language, expected):
    assert source_task_id(task_id, language) == expected


@pytest.mark.parametrize("langs", [["en"], ["en", "es", "en"]])
def test_a_side_by_side_sheet_needs_distinct_languages(langs):
    with pytest.raises(ValidationError):
        LocalizationSheetOptions(domain="retail", languages=langs)
