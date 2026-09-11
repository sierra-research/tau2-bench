# Copyright Sierra
"""Workbook builder: reads its own xlsx back to verify tabs, hidden columns,
dropdowns, and the blind-safety rule (rubric never shows non_native_ref)."""

from pathlib import Path

import pytest
from openpyxl import load_workbook

from tau2.annotation.nativeness_sheets import export_nativeness_sheets
from tau2.annotation.sheets.workbook import (
    build_and_publish,
    build_workbooks,
    readable_name,
)
from test_annotation.conftest import make_hi_results_dir


def _dropdown_formulas(ws):
    return {dv.formula1 for dv in ws.data_validations.dataValidation}


def _hidden_headers(ws):
    headers = [c.value for c in ws[1]]
    return {
        headers[i]
        for i in range(len(headers))
        if ws.column_dimensions[ws.cell(row=1, column=i + 1).column_letter].hidden
    }


@pytest.fixture
def both_family(tmp_path):
    """A real `--mode both` family exported from the two-sim hi fixture."""
    run_dir = make_hi_results_dir(tmp_path)
    out_stem = tmp_path / "hi_openai"
    export_nativeness_sheets([run_dir], "both", out_stem, reuse_existing=True)
    return tmp_path


def test_precision_workbook_dropdown_and_hidden(both_family):
    built = build_workbooks(both_family / "hi_openai_precision.csv")
    assert len(built) == 1
    out, mode_label = built[0]
    assert mode_label == "ADJUDICATE"
    wb = load_workbook(out)
    assert wb.sheetnames == ["Adjudicate"]
    ws = wb["Adjudicate"]
    assert ws.freeze_panes == "A2"
    # Verdict column carries the precision dropdown; trace ids are hidden.
    assert any("✅ yes" in f for f in _dropdown_formulas(ws))
    assert {"factor_id", "sim_id", "task_id", "language"} <= _hidden_headers(ws)
    # The readable column heads the sheet and the row data round-tripped.
    assert ws["A1"].value == "What we're checking"
    assert ws["A2"].value  # the nuance of the flagged factor


def test_cold_workbook_two_tabs_blind_safe_rubric(both_family):
    built = build_workbooks(both_family / "hi_openai.csv")
    out, mode_label = built[0]
    assert mode_label == "ANNOTATE"
    wb = load_workbook(out)
    assert wb.sheetnames == ["Conversations", "What each column means"]
    grid = wb["Conversations"]
    # Every factor column got the cold dropdown; trace ids hidden.
    assert any("no-opportunity" in f for f in _dropdown_formulas(grid))
    assert {"sim_id", "task_id", "language"} <= _hidden_headers(grid)
    # Rubric tab is blind-safe: nuance + the native standard only — the
    # non-native slip (which would prime labels) must NEVER appear.
    rubric = wb["What each column means"]
    rubric_text = "\n".join(
        str(c.value) for row in rubric.iter_rows() for c in row if c.value
    )
    from tau2.annotation.artifacts import read_artifact

    keys = read_artifact(both_family / "hi_openai.csv").rows("factors_key")
    for key in keys:
        assert key.nuance in rubric_text
        assert key.native_does in rubric_text
        assert key.non_native_ref not in rubric_text


def test_workbook_cells_never_become_formulas(tmp_path):
    """Content that starts with '=' must land as an inert text cell — never as
    an openpyxl formula (data_type 'f')."""
    from tau2.annotation.sheets.workbook import build_precision_workbook

    rows = [
        {
            "What we're checking": '=HYPERLINK("http://evil","click")',
            "Exact phrase flagged": "=1+1",
            "Real violation?": "",
        }
    ]
    out = build_precision_workbook(rows, tmp_path / "inert.xlsx")
    wb = load_workbook(out)
    ws = wb["Adjudicate"]
    for row in ws.iter_rows():
        for cell in row:
            assert cell.data_type != "f"
    # The '=' content survives verbatim as text.
    values = {c.value for row in ws.iter_rows() for c in row if c.value}
    assert '=HYPERLINK("http://evil","click")' in values
    assert "=1+1" in values


def test_workbook_mode_comes_from_manifest_not_headers(tmp_path):
    # A CSV with no manifest cannot build (mode detection is manifest-driven).
    orphan = tmp_path / "es_openai.csv"
    orphan.write_text("sim_id,X\ns1,\n")
    with pytest.raises(FileNotFoundError, match="manifest"):
        build_workbooks(orphan)


def test_sidecar_csv_is_not_a_workbook_input(both_family):
    with pytest.raises(ValueError, match="no workbook builder"):
        build_workbooks(both_family / "hi_openai_judge_sidecar.csv")


def test_readable_name():
    assert readable_name(Path("es_openai_precision.csv"), "ADJUDICATE") == (
        "es (openai) — ADJUDICATE.xlsx"
    )
    assert readable_name(Path("pt_openai.csv"), "ANNOTATE") == (
        "pt (openai) — ANNOTATE.xlsx"
    )


def test_publish_to_copies_with_readable_name(both_family):
    pub = both_family / "drive_mount"
    built = build_and_publish([both_family / "hi_openai_precision.csv"], publish_to=pub)
    assert built == 1
    assert (pub / "hi (openai) — ADJUDICATE.xlsx").exists()  # readable copy
    assert (both_family / "hi_openai_precision.xlsx").exists()  # original build
